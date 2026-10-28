"""把"手动/agent 查到的价格"导入今天的 sheet，比价规则和自动渠道完全一致。适用于 Allegro 等不适合脚本抓取的渠道。
用法:
  python app/import_manual.py --channel Allegro --csv data/allegro_input.csv [--table price_table.xlsx]
CSV 表头: sku,memory,color,price,status      (UTF-8，Excel 另存为 CSV UTF-8 即可)
  - 每个颜色一行；status 留空=有货，写 unavailable=缺货(页面变灰/到货通知/售罄)
  - 整款缺货: 写一行 color 留空、price 留空、status=unavailable
  - 没有这个型号的 offer: 不要写这一行(格子保持原样；想标注可以 status 写 no_offer，格子会写"该渠道无此商品")
前提: 今天的 sheet 已存在(先跑过 mistore 步骤)；price_log.xlsx 没有在 Excel 里打开。"""
import argparse, csv, datetime as dt, pathlib, sys
import openpyxl
import cell, hist, paths
from check_mistore import mem_key

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--channel", required=True); ap.add_argument("--csv", required=True)
    ap.add_argument("--table", default=str(paths.PRICE_TABLE)); a = ap.parse_args()
    if paths.EXCEL_LOCK.exists(): sys.exit("price_log.xlsx 正在被 Excel 打开。请先关闭 Excel(提示保存时选\"不保存\")，再重新运行。")
    ws_in = openpyxl.load_workbook(a.table, data_only=True).worksheets[0]
    hdr = {str(c.value).split("\n")[0].strip().lower(): c.column for c in ws_in[1] if c.value}
    rows = {}
    for r in range(2, ws_in.max_row + 1):
        sku, mem = ws_in.cell(r, hdr["sku"]).value, ws_in.cell(r, hdr["memory"]).value
        if sku and mem: rows[(str(sku).strip(), mem_key(mem))] = r
    today = dt.date.today(); master = openpyxl.load_workbook(paths.PRICE_LOG)
    if today.isoformat() not in master.sheetnames: sys.exit("今天的 sheet 还不存在。请先运行 mistore 步骤: python app/run_all.py --steps mistore")
    ws = master[today.isoformat()]
    ch = next((c for k, c in hdr.items() if k.startswith(a.channel.lower())), None)
    if ch is None: sys.exit(f"价格表里找不到渠道列: {a.channel}")
    groups = {}
    with open(a.csv, encoding="utf-8-sig", newline="") as f:
        for rec in csv.DictReader(f):
            rec = {(k or "").strip().lower(): (v or "").strip() for k, v in rec.items()}
            groups.setdefault((rec.get("sku", ""), mem_key(rec.get("memory", ""))), []).append(rec)
    log, missing = [], []
    for key, recs in groups.items():
        r = rows.get(key)
        if r is None: missing.append(key); continue
        promo = ws_in.cell(r, hdr["rrp after promotion"]).value
        if all(x.get("status", "").lower() == "no_offer" for x in recs): txt = "该渠道无此商品"
        else:
            items = [(x.get("color", ""), x.get("price") or None, x.get("status", "")) for x in recs if x.get("status", "").lower() != "no_offer"]
            txt = cell.summarize(items, promo)
            s, e = ws_in.cell(r, hdr["offer start date"]).value, ws_in.cell(r, hdr["offer end date"]).value
            s = s.date() if isinstance(s, dt.datetime) else s; e = e.date() if isinstance(e, dt.datetime) else e
            if isinstance(s, dt.date) and isinstance(e, dt.date) and not (s <= today <= e): txt = "[促销期外] " + txt
            log += [[today.isoformat(), a.channel, ws_in.cell(r, hdr["product name"]).value, ws_in.cell(r, hdr["memory"]).value, x.get("color", ""), x.get("price", ""), promo, x.get("status", "") or "available"] for x in recs]
        ws.cell(r, ch).value = txt; print(f"行{r} {ws_in.cell(r, hdr['product name']).value} {ws_in.cell(r, hdr['memory']).value}: {txt}")
    master.save(paths.PRICE_LOG); hist.append_rows(paths.HISTORY, log)
    for sku, mk in missing: print(f"⚠ CSV 里的 {sku} / {mk} 在价格表里找不到，已忽略")
    print(f"完成，已写入 {paths.PRICE_LOG.name} [{today.isoformat()}] 的 {a.channel} 列")

if __name__ == "__main__": main()
