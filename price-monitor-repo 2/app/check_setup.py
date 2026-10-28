"""配置体检: 对照价格表，列出每个型号在各渠道缺什么配置，告诉你(或 agent)下一步该补什么。只读，不联网。
用法: python app/check_setup.py [--table price_table.xlsx] [--json]
检查项: Mi-Store 对应关系(config.py MISTORE)、mi.com 标识(config.py MICOM)、EURO / OleOle / Allegro 链接(config/links.json，键 SKU|内存)。
新增机型后先跑它，再按提示补配置；补完再跑一次确认没有缺项。"""
import argparse, json, sys
import openpyxl
import config, paths

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--table", default=str(paths.PRICE_TABLE)); ap.add_argument("--json", action="store_true"); a = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    lf = paths.links_file(); links = json.load(open(lf, encoding="utf-8")) if lf.exists() else {}
    ws = openpyxl.load_workbook(a.table, data_only=True).worksheets[0]
    hdr = {str(c.value).split("\n")[0].strip().lower(): c.column for c in ws[1] if c.value}
    need = {"EURO": "EURO", "OleOle": "OleOle", "Allegro": "Allegro"}; out = []
    for r in range(2, ws.max_row + 1):
        sku, mem, name = (ws.cell(r, hdr[k]).value for k in ("sku", "memory", "product name"))
        if not (sku and mem and name): continue
        name = str(name).strip(); key = f"{sku}|{mem}" + ("|cover" if "cover" in name.lower() else ""); miss = []
        if name not in config.MISTORE: miss.append("Mi-Store: config.py 的 MISTORE 里没有这个型号名")
        if name not in config.MICOM: miss.append("mi.com: config.py 的 MICOM 里没有这个型号名")
        L = links.get(key, {})
        for ch in need:
            if not L.get(ch): miss.append(f"{ch}: links.json 缺少 \"{key}\" 的 {ch} 链接")
        out.append({"row": r, "sku": str(sku), "memory": str(mem), "name": name, "key": key, "missing": miss})
    bad = [o for o in out if o["missing"]]
    if a.json: print(json.dumps({"links_file": str(lf), "rows": len(out), "with_gaps": bad}, ensure_ascii=False, indent=1)); return
    print(f"链接库: {lf.name} | 价格表共 {len(out)} 行，{len(bad)} 行有配置缺项\n")
    for o in bad:
        print(f"行{o['row']} {o['name']} {o['memory']} ({o['sku']})")
        for m in o["missing"]: print("   -", m)
    if not bad: print("✔ 所有型号在所有渠道都已配置。")
    else: print("\n补法: Mi-Store/mi.com 改 app/config.py；EURO/OleOle/Allegro 链接用 python app/tools/add_link.py SKU 内存 渠道 链接 。详见 skills/price-monitor-operator/SKILL.md")

if __name__ == "__main__": main()
