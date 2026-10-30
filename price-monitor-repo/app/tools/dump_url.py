"""排查工具: 打开某行某站点(可指定颜色)的页面，保存网页并打印价格/库存字段。
用法: .venv/bin/python app/tools/dump_url.py 行号 EURO|OleOle [颜色关键词，如 fioletowy]   (需先开专用 Chrome)"""
import json, pathlib, re, sys
ROOT = pathlib.Path(__file__).resolve().parents[2]; sys.path.insert(0, str(ROOT / "app"))
import openpyxl, retail
from playwright.sync_api import sync_playwright
row, site = int(sys.argv[1]), sys.argv[2]; color = sys.argv[3].lower() if len(sys.argv) > 3 else None
ws = openpyxl.load_workbook(ROOT / "price_table.xlsx", data_only=True).worksheets[0]
url = json.load(open(__import__("paths").links_file(), encoding="utf-8"))[f"{ws.cell(row,1).value}|{ws.cell(row,2).value}"][site]
print("页面:", url, "| 表里促销价:", ws.cell(row, 5).value)
with sync_playwright() as p:
    pg = p.chromium.connect_over_cdp("http://localhost:9222").contexts[0].new_page()
    pg.goto(url, wait_until="domcontentloaded", timeout=45000); pg.wait_for_timeout(5000)
    html = pg.content()
    if color:
        slug = url.split("?")[0].rsplit("/", 1)[-1].replace(".bhtml", ""); origin = re.match(r"https?://[^/]+", url).group(0)
        for c, l, g in retail.color_group(html, slug) or []:
            if color in c.lower():
                pg.goto(f"{origin}/{g or url.split('/')[3]}/{l}.bhtml", wait_until="domcontentloaded", timeout=45000); pg.wait_for_timeout(5000); html = pg.content(); break
        else: print("没找到颜色", color)
    out = ROOT / "data" / "raw" / f"dump_{site}_row{row}{'_'+color if color else ''}.html"; out.write_text(html, encoding="utf-8"); print("已保存:", out.name)
    print("product_info:", retail.product_info(html)); print("delivery_status:", retail.delivery_status(html, (retail.product_info(html) or {}).get("availability", "")))
    i = html.find('"deliveryAvailability"'); print(html[i:i+500])
    for m in list(re.finditer(r"(?i)niedost\w+|powiadom\w*", html))[:8]: print("  文字:", re.sub(r"\s+", " ", html[max(0, m.start()-80):m.end()+60]))
    pg.close()
